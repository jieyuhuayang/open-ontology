"""Excel file parser — extracts columns, types, and metadata per sheet."""

from pathlib import Path

from app.agent.parsers.base import (
    BaseParser,
    ColumnInfo,
    ParseResult,
    is_audit_field,
    is_primary_key_candidate,
)
from app.agent.parsers.csv_parser import _infer_type, _MAX_INFER_ROWS, _SAMPLE_ROWS


class ExcelParser(BaseParser):
    async def parse(self, file_path: Path) -> ParseResult:
        import openpyxl

        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        all_columns: list[ColumnInfo] = []
        total_rows = 0
        sheet_info: list[dict] = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows_iter = ws.iter_rows(values_only=True)

            # Read header
            try:
                header_row = next(rows_iter)
            except StopIteration:
                continue

            header = [str(c) if c is not None else f"col_{i}" for i, c in enumerate(header_row)]

            # Read data rows
            data_rows: list[list[str]] = []
            for row in rows_iter:
                data_rows.append([str(c) if c is not None else "" for c in row])
                if len(data_rows) >= _MAX_INFER_ROWS:
                    break

            sheet_row_count = len(data_rows)
            total_rows += sheet_row_count

            columns: list[ColumnInfo] = []
            for i, col_name in enumerate(header):
                col_values = [row[i] for row in data_rows if i < len(row)]
                samples = [v for v in col_values[:_SAMPLE_ROWS] if v.strip()]

                columns.append(
                    ColumnInfo(
                        name=col_name.strip(),
                        inferred_type=_infer_type(col_values),
                        sample_values=samples,
                        is_primary_key_candidate=is_primary_key_candidate(col_name.strip()),
                        is_audit_field=is_audit_field(col_name.strip()),
                    )
                )

            all_columns.extend(columns)
            sheet_info.append(
                {"name": sheet_name, "columns": len(columns), "rows": sheet_row_count}
            )

        wb.close()

        return ParseResult(
            file_type="xlsx",
            columns=all_columns,
            row_count=total_rows,
            metadata={"sheets": sheet_info},
        )
